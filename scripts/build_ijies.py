"""Builds the IJIES Word version of the paper ON the official template.

The template's own sample paragraphs (title, authors, abstract, first- and
second-order headings, body text, captions, table, closing sections,
reference list) are deep-copied as prototypes and only their text is
replaced, so every page, section, style and direct-formatting setting of
"(Ver. 2026.10.1) IJIES_Format.docx" is kept exactly.

Content comes from paper/main.tex (single source of truth). Inline and
display math are converted to native Word equations (OMML) with pandoc;
tables are rebuilt from paper/tables/*.tex; references are formatted in
IJIES style from paper/refs.bib (registry-verified metadata).

    python scripts/build_ijies.py [template.docx]
        -> paper/ijies/EdgeIDS-KD_IJIES.docx
"""
import copy
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import docx
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"
OUT = PAPER / "ijies" / "EdgeIDS-KD_IJIES.docx"
TEMPLATE = Path(sys.argv[1]) if len(sys.argv) > 1 else PAPER / "ijies" / "IJIES_Format.docx"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS = {"w": W_NS, "m": M_NS}


def q(tag):
    p, t = tag.split(":")
    return f"{{{NS[p]}}}{t}"


# Twips: A4 11906 wide, margins 1080 -> text 9746; 2 columns, gap 432.
FULL_W, COL_W = 9746, (9746 - 432) // 2
FULL_CM, COL_CM = 17.19, 8.21

# ---------------------------------------------------------------- content
ABSTRACT = (
    "Intrusion detectors deployed on Internet of Things devices must be compressed to fit "
    "kilobyte-scale memory, yet compressed models are usually judged by overall accuracy, which "
    "can hide the loss of rare attacks. We study knowledge distillation and post-training integer "
    "quantization for IoT intrusion detection on five public datasets after a leakage audit that "
    "removes duplicated and identifier-dependent records. Across seven student sizes, seven "
    "training objectives and three seeds, plain cross-entropy students lose almost all "
    "rare-attack recall at the smallest sizes while their macro F1 remains moderate. Small "
    "students are rescued by imbalance-aware losses rather than by distillation itself, whereas "
    "distilling from a logit-adjusted teacher raises rare-attack recall on every dataset at "
    "little accuracy cost. Most of the accuracy lost to INT8 quantization originates in the input "
    "layer and is recovered by keeping that layer in floating point. Every student fits a 400 KB "
    "microcontroller budget, whereas neither teacher does.")
KEYWORDS = ("Intrusion detection, Internet of things, Knowledge distillation, Quantization, "
            "Class imbalance, Edge computing")

# The objectives bullet list of main.tex becomes numbered equations (the
# template requires properly formatted, numbered equations).
OBJECTIVES = r"""
\textbf{CE and WCE.} Cross-entropy (CE) and its inverse-frequency weighted form (WCE) are
\begin{equation}
\mathcal{L}_{\mathrm{WCE}} = -\frac{\sum_{i\in B} w_{y_i}\log\sigma(s_i)_{y_i}}{\sum_{i\in B} w_{y_i}}
\label{eq:wce}
\end{equation}
with $w_c=1/(C\pi_c)$ normalised to mean 1 over the $C$ classes; CE is the case $w_c=1$.

\textbf{LA.} Logit-adjusted CE~\citep{menon2021logit} trains on prior-shifted logits and predicts with the raw $s$ ($\tau=1$):
\begin{equation}
\mathcal{L}_{\mathrm{LA}} = \mathrm{CE}\big(s+\tau\log\pi,\ y\big)
\label{eq:la}
\end{equation}

\textbf{KD and KD+WCE.} Distillation~\citep{hinton2015distilling} adds a softened Kullback--Leibler term to a hard-label loss:
\begin{equation}
\mathcal{L}_{\mathrm{KD}} = \alpha\,\mathcal{L}_{\mathrm{soft}} + (1-\alpha)\,\mathcal{L}_{\mathrm{hard}}
\label{eq:kd}
\end{equation}
\begin{equation}
\mathcal{L}_{\mathrm{soft}} = T^2\,\mathrm{KL}\big(\sigma(t/T)\,\|\,\sigma(s/T)\big)
\label{eq:soft}
\end{equation}
with $\mathcal{L}_{\mathrm{hard}}$ equal to CE for KD and to WCE for KD+WCE.

\textbf{RAP-KD (ours).} The KL term is weighted per sample by $\tilde w_{y_i}\propto w_{y_i}^{\beta}$ ($\beta=0.5$, normalised to mean 1), and the hard term is the LA loss:
\begin{equation}
\mathcal{L}_{\mathrm{RAP}} = \alpha T^2\,\frac{\sum_{i\in B}\tilde w_{y_i}\,\mathrm{KL}_i}{\sum_{i\in B}\tilde w_{y_i}} + (1-\alpha)\,\mathcal{L}_{\mathrm{LA}}
\label{eq:rap}
\end{equation}

\textbf{LA-KD (ours).} Both terms of Eq.~(\ref{eq:kd}) are computed on the prior-shifted logits $s+\tau\log\pi$.

\textbf{Balanced teacher (BT).} The teacher is retrained with Eq.~(\ref{eq:la}), and the student is distilled from it with KD, KD+WCE, or BT-KD, a variant of ours that uses Eq.~(\ref{eq:kd}) with $\mathcal{L}_{\mathrm{hard}}=\mathcal{L}_{\mathrm{LA}}$.
"""

CLOSING = {
    "Data and Code Availability": None,          # taken from main.tex
    "Conflicts of Interest": "The authors declare no conflict of interest.",
    "Author Contributions": ("\\ph{Conceptualization, XXX and YYY; methodology, XXX; software, XXX; "
                             "validation, XXX and YYY; formal analysis, XXX; investigation, XXX; "
                             "data curation, XXX; writing---original draft preparation, XXX; "
                             "writing---review and editing, XXX; visualization, XXX; supervision, YYY.}"),
    "Acknowledgments": "\\ph{[Funding and support, if any; otherwise remove this section.]}",
}


# ------------------------------------------------------------ tex utils
def find_group(s, i):
    """s[i] == '{' -> (content, index after the matching '}')."""
    assert s[i] == "{", s[i:i + 30]
    depth, j = 0, i
    while j < len(s):
        if s[j] == "\\":
            j += 2
            continue
        if s[j] == "{":
            depth += 1
        elif s[j] == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    raise ValueError("unbalanced braces: " + s[i:i + 60])


def strip_comments(s):
    return re.sub(r"(?<!\\)%.*", "", s)


# ---------------------------------------------------------- bibliography
def parse_bib(path):
    txt = path.read_text(encoding="utf-8")
    entries = {}
    for m in re.finditer(r"@(\w+)\{([^,]+),", txt):
        kind, key = m.group(1), m.group(2).strip()
        i = txt.index("{", m.start())
        body, _ = find_group(txt, i)
        fields = {}
        for fm in re.finditer(r"(\w+)\s*=\s*\{", body):
            val, _ = find_group(body, fm.end() - 1)
            fields[fm.group(1).lower()] = val
        entries[key] = (kind, fields)
    return entries


def plain(s):
    s = re.sub(r"\\href\{[^}]*\}\{([^}]*)\}", r"\1", s)
    s = s.replace("\\&", "&").replace("\\_", "_").replace("\\%", "%").replace("--", "-")
    s = s.replace("{", "").replace("}", "").replace("~", " ")
    return re.sub(r"\s+", " ", s).strip()


def initials(given):
    parts = []
    for tok in given.replace(".", ". ").split():
        sub = [p[0] + "." for p in tok.split("-") if p and p[0].isalpha()]
        if sub:
            parts.append("-".join(sub))
    return " ".join(parts)


def authors_ijies(field):
    names = []
    for a in [x.strip() for x in field.split(" and ") if x.strip()]:
        if "," in a:
            fam, giv = [p.strip() for p in a.split(",", 1)]
        else:
            toks = a.split()
            fam, giv = toks[-1], " ".join(toks[:-1])
        names.append(f"{initials(plain(giv))} {plain(fam)}".strip())
    if len(names) == 1:
        return names[0]
    if len(names) == 2:
        return f"{names[0]} and {names[1]}"
    return ", ".join(names[:-1]) + ", and " + names[-1]


def surname(field):
    first = field.split(" and ")[0].strip()
    fam = first.split(",")[0] if "," in first else first.split()[-1]
    n = len([x for x in field.split(" and ") if x.strip()])
    fam = plain(fam)
    if n == 1:
        return fam
    if n == 2:
        second = field.split(" and ")[1].strip()
        fam2 = plain(second.split(",")[0] if "," in second else second.split()[-1])
        return f"{fam} and {fam2}"
    return f"{fam} et al."


def reference_runs(kind, f):
    """IJIES style: A. B. Author, C. Author, and D. Author, "Title",
    Journal, Vol. x, No. y, pp. a-b, year.  /  In: Proc. of <conf>, pp., year."""
    runs = [(authors_ijies(f.get("author", "")) + ", ", {}),
            ("\u201c" + plain(f["title"]) + "\u201d, ", {})]
    note = f.get("note", "")
    doi = f.get("doi")
    pub_note = plain(re.sub(r"doi:.*", "", note)).rstrip(". ")
    if kind == "inproceedings":
        runs.append(("In: ", {}))
        runs.append(("Proc. of " + plain(f.get("booktitle", "")), {"italic": True}))
    else:
        runs.append((plain(f.get("journal", "")), {"italic": True}))
    tail = []
    if f.get("volume"):
        tail.append("Vol. " + plain(f["volume"]))
    if f.get("number"):
        tail.append("No. " + plain(f["number"]))
    if f.get("pages"):
        tail.append("pp. " + plain(f["pages"]))
    if pub_note:
        tail.append(pub_note[0].lower() + pub_note[1:])
    tail.append(f.get("year", ""))
    text = ", " + ", ".join(t for t in tail if t)
    if doi:
        text += f", doi: {doi}"
    runs.append((text + ".", {}))
    return runs


# --------------------------------------------------------------- math
class MathConverter:
    """Batch LaTeX -> OMML with pandoc (one docx round trip for all snippets)."""

    def __init__(self):
        self.snippets, self.cache = [], {}

    def want(self, tex):
        tex = tex.strip()
        if tex not in self.cache and tex not in self.snippets:
            self.snippets.append(tex)
        return tex

    def run(self):
        if not self.snippets:
            return
        with tempfile.TemporaryDirectory() as tmp:
            md = Path(tmp) / "m.md"
            md.write_text("\n\n".join(f"${s}$" for s in self.snippets), encoding="utf-8")
            out = Path(tmp) / "m.docx"
            subprocess.run(["pandoc", str(md), "-f", "markdown", "-o", str(out)], check=True)
            xml = zipfile.ZipFile(out).read("word/document.xml")
        root = etree.fromstring(xml)
        paras = root.findall(".//w:body/w:p", NS)
        if len(paras) != len(self.snippets):
            raise RuntimeError(f"pandoc returned {len(paras)} paragraphs for {len(self.snippets)} snippets")
        for s, p in zip(self.snippets, paras):
            oms = p.findall(".//m:oMath", NS)
            if len(oms) != 1:
                raise RuntimeError(f"math not converted: {s!r}")
            self.cache[s] = oms[0]
        self.snippets = []

    def omml(self, tex):
        return copy.deepcopy(self.cache[tex.strip()])


SIMPLE_MATH = {r"\times": "\u00d7", r"\pm": "\u00b1", r"\mu": "\u00b5", r"\geq": "\u2265",
               r"\leq": "\u2264", r"\alpha": "\u03b1", r"\beta": "\u03b2", r"\tau": "\u03c4",
               r"\rho": "\u03c1", r"\pi": "\u03c0", r"\sim": "~", r"\approx": "\u2248",
               r"\,": "\u2009", r"\!": "", "~": "\u00a0"}


def simple_math_text(tex):
    """Plain-text rendering for trivial math (numbers, signs, single symbols)."""
    t = tex.strip()
    if re.fullmatch(r"[\d\s.,+\-=<>%]*(\\(times|pm|mu|geq|leq|,|alpha|beta|tau|rho|pi|approx|sim))?[\d\s.,+\-=<>%a-zA-Z]*", t) \
            and not re.search(r"[a-zA-Z]{2,}", re.sub(r"\\\w+", "", t)):
        for k, v in SIMPLE_MATH.items():
            t = t.replace(k, v)
        t = t.replace("-", "\u2212") if re.fullmatch(r"-[\d.]+", t) else t
        return t.replace("{", "").replace("}", "")
    return None


# ---------------------------------------------------------- inline TeX
class Inline:
    """LaTeX inline -> list of segments: ('t', text, fmt) | ('m', tex)."""

    def __init__(self, ctx):
        self.ctx = ctx

    def parse(self, s, fmt=None):
        fmt = dict(fmt or {})
        out, buf, i = [], [], 0

        def flush():
            if buf:
                out.append(("t", "".join(buf), dict(fmt)))
                buf.clear()

        while i < len(s):
            c = s[i]
            if c == "$":
                j = s.index("$", i + 1)
                flush()
                tex = s[i + 1:j]
                simple = simple_math_text(tex)
                if simple is not None:
                    f2 = dict(fmt)
                    if re.search(r"[A-Za-z]", simple):      # variables are italic
                        f2["italic"] = True
                    out.append(("t", simple, f2))
                else:
                    out.append(("m", self.ctx.math.want(tex)))
                i = j + 1
                continue
            if c == "\\":
                m = re.match(r"\\([A-Za-z]+)\*?", s[i:])
                if not m:
                    sym = s[i + 1:i + 2]
                    buf.append({",": "\u2009", " ": " ", "%": "%", "_": "_", "&": "&", "#": "#",
                                "{": "{", "}": "}", "-": "", "\\": " "}.get(sym, sym))
                    i += 2
                    continue
                cmd, i = m.group(1), i + m.end()
                if cmd in ("textbf", "emph", "textit", "texttt", "ph", "textsc"):
                    arg, i = find_group(s, i)
                    flush()
                    f2 = dict(fmt)
                    f2.update({"textbf": {"bold": True}, "emph": {"italic": True},
                               "textit": {"italic": True}, "texttt": {"mono": True},
                               "ph": {"red": True}, "textsc": {}}[cmd])
                    out.extend(self.parse(arg, f2))
                    continue
                if cmd in ("citep", "cite", "citet"):
                    arg, i = find_group(s, i)
                    flush()
                    keys = [k.strip() for k in arg.split(",")]
                    if cmd == "citet":
                        out.append(("t", self.ctx.cite_names(keys) + " ", dict(fmt)))
                    out.append(("t", self.ctx.cite_numbers(keys), dict(fmt)))
                    continue
                if cmd == "ref":
                    arg, i = find_group(s, i)
                    buf.append(self.ctx.ref(arg))
                    continue
                if cmd in ("label",):
                    _, i = find_group(s, i)
                    continue
                if cmd in ("ldots", "dots"):
                    buf.append("\u2026")
                    continue
                if cmd == "tiny" or cmd == "scriptsize" or cmd == "small":
                    continue
                buf.append({"times": "\u00d7", "pm": "\u00b1", "mu": "\u00b5"}.get(cmd, ""))
                continue
            if c in "{}":
                i += 1
                continue
            if c == "~":
                buf.append("\u00a0")
                i += 1
                continue
            if s.startswith("---", i):
                buf.append("\u2014")
                i += 3
                continue
            if s.startswith("--", i):
                buf.append("\u2013")
                i += 2
                continue
            if s.startswith("``", i):
                buf.append("\u201c")
                i += 2
                continue
            if s.startswith("''", i):
                buf.append("\u201d")
                i += 2
                continue
            if c == "\n":
                buf.append(" ")
                i += 1
                continue
            buf.append(c)
            i += 1
        flush()
        # collapse whitespace inside text segments
        cleaned = []
        for seg in out:
            if seg[0] == "t":
                t = re.sub(r"[ \t]+", " ", seg[1])
                cleaned.append(("t", t, seg[2]))
            else:
                cleaned.append(seg)
        return cleaned


# --------------------------------------------------------- document model
class Builder:
    def __init__(self, template):
        self.doc = docx.Document(template)
        self.body = self.doc.element.body
        P = self.doc.paragraphs
        self.proto = {
            "h1": copy.deepcopy(P[15]._p), "body": copy.deepcopy(P[16]._p),
            "figpara": copy.deepcopy(P[37]._p), "caption": copy.deepcopy(P[42]._p),
            "tabtitle": copy.deepcopy(P[44]._p), "blank": copy.deepcopy(P[46]._p),
            "chapter": copy.deepcopy(P[56]._p), "textbody": copy.deepcopy(P[57]._p),
            "refhead": copy.deepcopy(P[63]._p), "refitem": copy.deepcopy(P[67]._p),
            "table": copy.deepcopy(self.doc.tables[0]._tbl),
        }
        self.title_block = P[:15]
        # section properties: 1-column title section (P14), 2-column body, 1-column full width
        self.sect_two = copy.deepcopy(self._sect_of_section(1))
        self.sect_one = copy.deepcopy(self._sect_of_section(2))
        for s in (self.sect_two, self.sect_one):
            t = s.find(q("w:type"))
            if t is None:
                t = etree.SubElement(s, q("w:type"))
                s.insert(0, t)
            t.set(q("w:val"), "continuous")
        self.math = MathConverter()
        self.inline = Inline(self)
        self.cite_order, self.labels = {}, {}
        self.bib = parse_bib(PAPER / "refs.bib")
        self.ops = []          # deferred operations (math needs one pandoc pass)

    def _sect_of_section(self, k):
        sects = [p.find(q("w:pPr") + "/" + q("w:sectPr")) for p in self.body.findall(q("w:p"))]
        sects = [s for s in sects if s is not None] + [self.body.find(q("w:sectPr"))]
        return sects[k]

    # ----- citations / references
    def cite_numbers(self, keys):
        nums = []
        for k in keys:
            if k not in self.bib:
                raise KeyError(f"unknown citation {k}")
            self.cite_order.setdefault(k, len(self.cite_order) + 1)
            nums.append(self.cite_order[k])
        nums = sorted(set(nums))
        parts, start = [], None
        for i, n in enumerate(nums):
            if start is None:
                start = n
            if i + 1 == len(nums) or nums[i + 1] != n + 1:
                parts.append(str(start) if start == n else f"{start}\u2013{n}")
                start = None
        return "[" + ", ".join(parts) + "]"

    def cite_names(self, keys):
        return " and ".join(surname(self.bib[k][1].get("author", "")) for k in keys)

    def ref(self, label):
        if label not in self.labels:
            raise KeyError(f"unresolved \\ref{{{label}}}")
        return self.labels[label]

    # ----- element factories (prototype clones)
    def _clone(self, key):
        el = copy.deepcopy(self.proto[key])
        for r in el.findall(q("w:r")) + el.findall(q("w:hyperlink")):
            el.remove(r)
        for extra in el.findall(".//" + q("m:oMath")):
            extra.getparent().remove(extra)
        return el

    def append(self, el):
        self.body.find(q("w:sectPr")).addprevious(el)
        return el

    def _run_template(self, p_el):
        """rPr to use for runs in this paragraph: the prototype's pPr/rPr."""
        rpr = p_el.find(q("w:pPr") + "/" + q("w:rPr"))
        r = etree.SubElement(etree.Element(q("w:p")), q("w:r"))
        if rpr is not None:
            new = copy.deepcopy(rpr)
            for bad in ("vertAlign", "color"):
                for e in new.findall(q("w:" + bad)):
                    new.remove(e)
            r.append(new)
        return r

    def add_runs(self, p_el, segments, size=None, base=None):
        for seg in segments:
            if seg[0] == "br":
                r = self._run_template(p_el)
                etree.SubElement(r, q("w:br"))
                p_el.append(r)
                continue
            if seg[0] == "m":
                self.ops.append(("math", p_el, seg[1]))
                p_el.append(etree.Element(q("w:r")))         # placeholder, replaced later
                p_el[-1].set("mathph", seg[1])
                continue
            _, text, fmt = seg
            if not text:
                continue
            r = self._run_template(p_el)
            rpr = r.find(q("w:rPr"))
            if rpr is None:
                rpr = etree.SubElement(r, q("w:rPr"))
                r.insert(0, rpr)
            for e in rpr.findall(q("w:b")):
                rpr.remove(e)
            if fmt.get("bold") or (base or {}).get("bold"):
                rpr.append(etree.Element(q("w:b")))
            if fmt.get("italic"):
                rpr.append(etree.Element(q("w:i")))
            if fmt.get("mono"):
                f = etree.Element(q("w:rFonts"))
                for a in ("ascii", "hAnsi", "cs"):
                    f.set(q("w:" + a), "Courier New")
                for e in rpr.findall(q("w:rFonts")):
                    rpr.remove(e)
                rpr.insert(0, f)
            if fmt.get("red"):
                c = etree.Element(q("w:color"))
                c.set(q("w:val"), "FF0000")
                rpr.append(c)
            if fmt.get("sup"):
                v = etree.Element(q("w:vertAlign"))
                v.set(q("w:val"), "superscript")
                rpr.append(v)
            if fmt.get("sub"):
                v = etree.Element(q("w:vertAlign"))
                v.set(q("w:val"), "subscript")
                rpr.append(v)
            if size:
                for e in rpr.findall(q("w:sz")) + rpr.findall(q("w:szCs")):
                    rpr.remove(e)
                for t in ("sz", "szCs"):
                    e = etree.Element(q("w:" + t))
                    e.set(q("w:val"), str(int(size * 2)))
                    rpr.append(e)
            t = etree.SubElement(r, q("w:t"))
            t.text = text
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            p_el.append(r)

    def para(self, key, tex=None, segments=None, size=None, base=None):
        p = self._clone(key)
        self.add_runs(p, segments if segments is not None else self.inline.parse(tex), size, base)
        return self.append(p)

    def heading(self, level, text):
        p = self._clone("h1")
        if level == 2:
            p.find(q("w:pPr") + "/" + q("w:numPr") + "/" + q("w:ilvl")).set(q("w:val"), "1")
            sp = p.find(q("w:pPr") + "/" + q("w:spacing"))
            for a in ("beforeLines", "afterLines"):
                sp.attrib.pop(q("w:" + a), None)
            sp.set(q("w:before"), "160")
            sp.set(q("w:after"), "80")
        self.add_runs(p, self.inline.parse(text), size=11 if level == 2 else None, base={"bold": True})
        return self.append(p)

    def section_break(self, two_columns_before):
        """End the current section; the sectPr describes the section that ends."""
        p = self._clone("blank")
        ppr = p.find(q("w:pPr"))
        ppr.append(copy.deepcopy(self.sect_two if two_columns_before else self.sect_one))
        # keep the break paragraph minimal
        for e in ppr.findall(q("w:spacing")):
            ppr.remove(e)
        sp = etree.SubElement(ppr, q("w:spacing"))
        sp.set(q("w:after"), "0")
        sp.set(q("w:line"), "20")
        sp.set(q("w:lineRule"), "exact")
        ppr.remove(sp)
        ppr.insert(1, sp)
        self.append(p)

    def picture(self, path, width_cm):
        p = self._clone("figpara")
        self.append(p)
        Paragraph(p, self.doc._body).add_run().add_picture(str(path), width=Cm(width_cm))
        return p

    def equation(self, tex, number):
        p = self._clone("body")
        ppr = p.find(q("w:pPr"))
        for e in ppr.findall(q("w:ind")):
            ppr.remove(e)
        ind = etree.SubElement(ppr, q("w:ind"))
        ind.set(q("w:left"), "284")                     # 5 mm, as the template requires
        tabs = etree.Element(q("w:tabs"))
        tab = etree.SubElement(tabs, q("w:tab"))
        tab.set(q("w:val"), "right")
        tab.set(q("w:pos"), str(COL_W - 284))
        ppr.insert(1, tabs)
        jc = ppr.find(q("w:jc"))
        if jc is not None:
            jc.set(q("w:val"), "left")
        self.add_runs(p, [("m", self.math.want(tex))])
        r = self._run_template(p)
        etree.SubElement(r, q("w:tab"))
        t = etree.SubElement(r, q("w:t"))
        t.text = f"({number})"
        p.append(r)
        self.append(self._clone("blank"))
        self.append(p)
        self.append(self._clone("blank"))

    # ----- table from generated LaTeX
    def table(self, tex_path, width):
        src = strip_comments(tex_path.read_text(encoding="utf-8"))
        src = src[src.index("\\toprule"):src.rindex("\\bottomrule")]
        lines = [x for x in re.split(r"\\\\", src)]
        rows, bold_rows = [], 0
        for chunk in lines:
            chunk = re.sub(r"\\(toprule|midrule|bottomrule)", "", chunk)
            chunk = re.sub(r"\\cmidrule\([^)]*\)\{[^}]*\}", "", chunk).strip()
            if not chunk:
                continue
            cells = []
            for c in [x.strip() for x in re.split(r"(?<!\\)&", chunk)]:
                span = 1
                m = re.match(r"\\multicolumn\{(\d+)\}\{[^}]*\}\{(.*)\}$", c, re.S)
                if m:
                    span, c = int(m.group(1)), m.group(2)
                cells.append((span, c))
            rows.append(cells)
        n_cols = sum(s for s, _ in rows[-1])
        header_rows = sum(1 for r in rows if any("multicolumn" in c for _, c in r)) + 1
        has_pm = "\\pm" in src

        def cell_segments(content, col):
            segs = self.inline.parse(content)
            # tables use plain-text math only (10 pt Times throughout)
            segs = [s if s[0] == "t" else ("t", self.table_math(s[1]), {}) for s in segs]
            out = []
            for s in segs:
                t = s[1].replace(" ", " ")
                # no line breaks at dashes or hyphens inside a value ("12.1–13.2", "MLP-256")
                t = re.sub(r"(?<=\w)([–\-])(?=\w)", "⁠\\1⁠", t)
                if has_pm:
                    # mean on the first line, "±s.d." on the second; "(params)" likewise
                    pieces = re.split(r" ?(?=±)|\s(?=\()", t)
                    for k, piece in enumerate(pieces):
                        if k:
                            out.append(("br",))
                        out.append(("t", piece, s[2]))
                else:
                    out.append(("t", t, s[2]))
            return out

        def seg_len(segs, header):
            lines = [[]]
            for s in segs:
                if s[0] == "br":
                    lines.append([])
                else:
                    lines[-1].append(s[1].replace("⁠", ""))
            texts = ["".join(x) for x in lines]
            if header:   # header cells may wrap between words
                return max((len(w) for t in texts for w in t.split()), default=1)
            return max(len(t) for t in texts)

        cell_segs = [[cell_segments(c, i) for i, (_, c) in enumerate(cells)] for cells in rows]
        # column widths proportional to content (equal widths made numbers wrap)
        need = [3] * n_cols
        for ri, cells in enumerate(rows):
            col = 0
            for ci, (span, _) in enumerate(cells):
                if span == 1:
                    need[col] = max(need[col], seg_len(cell_segs[ri][ci], ri < header_rows))
                col += span
        widths = [max(1, int(width * n / sum(need))) for n in need]
        tbl = copy.deepcopy(self.proto["table"])
        for tr in tbl.findall(q("w:tr")):
            tbl.remove(tr)
        grid = tbl.find(q("w:tblGrid"))
        for g in list(grid):
            grid.remove(g)
        for wdt in widths:
            g = etree.SubElement(grid, q("w:gridCol"))
            g.set(q("w:w"), str(wdt))
        tblpr = tbl.find(q("w:tblPr"))
        tw = tblpr.find(q("w:tblW"))
        tw.set(q("w:w"), str(sum(widths)))
        tw.set(q("w:type"), "dxa")
        lay = etree.SubElement(tblpr, q("w:tblLayout"))
        lay.set(q("w:type"), "autofit")
        jc = etree.SubElement(tblpr, q("w:jc"))
        jc.set(q("w:val"), "center")
        tblpr.find(q("w:tblInd")).set(q("w:w"), "0")
        proto_tr = self.proto["table"].findall(q("w:tr"))
        for ri, cells in enumerate(rows):
            tr = copy.deepcopy(proto_tr[0 if ri < header_rows else 1])
            for tc in tr.findall(q("w:tc")):
                tr.remove(tc)
            proto_tc = proto_tr[0 if ri < header_rows else 1].find(q("w:tc"))
            col = 0
            for ci, (span, content) in enumerate(cells):
                tc = copy.deepcopy(proto_tc)
                tcpr = tc.find(q("w:tcPr"))
                tcpr.find(q("w:tcW")).set(q("w:w"), str(sum(widths[col:col + span])))
                col += span
                if span > 1:
                    gs = etree.SubElement(tcpr, q("w:gridSpan"))
                    gs.set(q("w:val"), str(span))
                    tcpr.insert(1, gs)
                p = tc.find(q("w:p"))
                for r in p.findall(q("w:r")):
                    p.remove(r)
                sp = p.find(q("w:pPr") + "/" + q("w:spacing"))
                if sp is None:
                    sp = etree.SubElement(p.find(q("w:pPr")), q("w:spacing"))
                sp.set(q("w:after"), "0")
                sp.set(q("w:line"), "240")
                sp.set(q("w:lineRule"), "auto")
                self.add_runs(p, cell_segs[ri][ci], size=10, base={"bold": ri < header_rows})
                tr.append(tc)
            tbl.append(tr)
        return self.append(tbl)

    @staticmethod
    def table_math(tex):
        t = tex
        t = re.sub(r"_\{\\mathrm\{([^}]*)\}\}", r"_\1", t)
        t = re.sub(r"\\mathrm\{([^}]*)\}", r"\1", t)
        for k, v in SIMPLE_MATH.items():
            t = t.replace(k, v)
        return t.replace("{", "").replace("}", "").replace("$", "")

    # ----- algorithm (boxed pseudocode, 10 pt, numbered lines)
    def algorithm(self, caption, label_num, body):
        lines = self.parse_algorithmic(body)
        tbl = copy.deepcopy(self.proto["table"])
        for tr in tbl.findall(q("w:tr")):
            tbl.remove(tr)
        grid = tbl.find(q("w:tblGrid"))
        for g in list(grid):
            grid.remove(g)
        g = etree.SubElement(grid, q("w:gridCol"))
        g.set(q("w:w"), str(COL_W))
        tblpr = tbl.find(q("w:tblPr"))
        tblpr.find(q("w:tblW")).set(q("w:w"), str(COL_W))
        tblpr.find(q("w:tblW")).set(q("w:type"), "dxa")
        tblpr.find(q("w:tblInd")).set(q("w:w"), "0")
        borders = etree.SubElement(tblpr, q("w:tblBorders"))
        for side in ("top", "bottom"):
            b = etree.SubElement(borders, q("w:" + side))
            for k, v in (("val", "single"), ("sz", "8"), ("space", "0"), ("color", "000000")):
                b.set(q("w:" + k), v)
        for side in ("left", "right", "insideH", "insideV"):
            b = etree.SubElement(borders, q("w:" + side))
            b.set(q("w:val"), "nil")
        proto_tr = self.proto["table"].findall(q("w:tr"))[1]
        tr = copy.deepcopy(proto_tr)
        for tc in tr.findall(q("w:tc")):
            tr.remove(tc)
        tc = copy.deepcopy(proto_tr.find(q("w:tc")))
        tc.find(q("w:tcPr")).find(q("w:tcW")).set(q("w:w"), str(COL_W))
        for p in tc.findall(q("w:p")):
            tc.remove(p)
        proto_p = proto_tr.find(q("w:tc")).find(q("w:p"))

        def line(segs, indent=0, number=None, hanging=False, border_bottom=False):
            p = copy.deepcopy(proto_p)
            for r in p.findall(q("w:r")):
                p.remove(r)
            ppr = p.find(q("w:pPr"))
            ppr.find(q("w:jc")).set(q("w:val"), "left")
            sp = etree.Element(q("w:spacing"))
            sp.set(q("w:after"), "0")
            sp.set(q("w:line"), "240")
            sp.set(q("w:lineRule"), "auto")
            ppr.insert(0, sp)
            # line numbers stay flush left; only the text is indented by level
            ind = etree.Element(q("w:ind"))
            ind.set(q("w:left"), str(340 + 240 * indent))
            ind.set(q("w:hanging"), str(340 + 240 * indent))
            ppr.insert(1, ind)
            if border_bottom:
                pb = etree.Element(q("w:pBdr"))
                b = etree.SubElement(pb, q("w:bottom"))
                for k, v in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", "000000")):
                    b.set(q("w:" + k), v)
                ppr.insert(0, pb)
            lead = [("t", f"{number}:\t" if number else "\t", {})] if not hanging else []
            self.add_runs(p, lead + segs, size=10)
            tc.append(p)

        line(self.inline.parse(f"\\textbf{{Algorithm {label_num}.}} {caption}"), hanging=True,
             border_bottom=True)
        n = 0
        for kind, indent, segs in lines:
            if kind == "head":
                line(segs, indent=0, hanging=True)
            else:
                n += 1
                line(segs, indent=indent, number=n)
        tr.append(tc)
        tbl.append(tr)
        self.append(tbl)
        self.append(self._clone("blank"))

    def parse_algorithmic(self, body):
        body = strip_comments(body)
        toks = re.split(r"(\\(?:Require|Ensure|State|Statex|For|EndFor|If|ElsIf|Else|EndIf|While|EndWhile|Return)\b)", body)
        out, level, i = [], 0, 1
        kw = lambda w: ("t", w, {"bold": True})  # noqa: E731

        def split_comment(text):
            m = re.search(r"\\Comment\{", text)
            if not m:
                return text, None
            c, _ = find_group(text, m.end() - 1)
            return text[:m.start()], c

        while i < len(toks):
            cmd, arg = toks[i], toks[i + 1]
            i += 2
            if cmd in ("\\For", "\\If", "\\ElsIf", "\\While"):
                a = arg.lstrip()
                cond, k = find_group(a, 0)
                rest_text, com = split_comment(a[k:])
                word = {"\\For": "for", "\\If": "if", "\\ElsIf": "else if", "\\While": "while"}[cmd]
                if cmd == "\\ElsIf":
                    level -= 1
                segs = [kw(word + " ")] + self.inline.parse(cond) + [kw(" do" if word in ("for", "while") else " then")]
                if com:
                    segs += [("t", "  \u25b7 ", {})] + self.inline.parse(com)
                out.append(("line", level, segs))
                level += 1
                # a \State on the same source line is already a separate token
                continue
            if cmd == "\\Else":
                out.append(("line", level - 1, [kw("else")]))
                continue
            if cmd in ("\\EndFor", "\\EndIf", "\\EndWhile"):
                level -= 1
                out.append(("line", level, [kw("end " + cmd[4:].lower().replace("if", "if"))]))
                continue
            text, com = split_comment(arg)
            text = text.strip()
            if cmd in ("\\Require", "\\Ensure"):
                segs = [kw("Input: " if cmd == "\\Require" else "Output: ")] + self.inline.parse(text)
                out.append(("head", 0, segs))
                continue
            if cmd == "\\Return":
                segs = [kw("return ")] + self.inline.parse(text)
            elif cmd == "\\Statex":
                segs = self.inline.parse(com or text)
                out.append(("head", level, [("t", "\u25b7 ", {})] + segs))
                continue
            else:
                if not text and not com:
                    continue        # "\State \Return x" splits into an empty State
                # inline control words used inside a State, e.g. "\textbf{break}"
                segs = self.inline.parse(text)
            if com:
                segs += [("t", "  \u25b7 ", {})] + self.inline.parse(com)
            out.append(("line", level, segs))
        return out

    # ----- finalisation
    def resolve_math(self):
        self.math.run()
        for p in self.body.iter(q("w:p")):
            for r in list(p.findall(q("w:r"))):
                tex = r.get("mathph")
                if tex is not None:
                    r.addprevious(self.math.omml(tex))
                    p.remove(r)


# ------------------------------------------------------------ the paper
def numbering_pass(src):
    """Labels -> printed numbers, following the IJIES conventions."""
    labels, sec, sub = {}, 0, 0
    fig = tab = alg = eq = 0
    in_subsec_paragraph_section = {}
    for m in re.finditer(r"\\(section|subsection|paragraph)\*?\{|\\begin\{(figure\*?|table\*?|algorithm|equation)\}|\\label\{([^}]+)\}", src):
        if m.group(1) == "section":
            if src[m.start():m.start() + 9] == "\\section*":
                last = None
                continue
            sec, sub = sec + 1, 0
            last = ("sec", f"{sec}")
            in_subsec_paragraph_section[sec] = False
        elif m.group(1) == "subsection":
            sub += 1
            in_subsec_paragraph_section[sec] = True
            last = ("sub", f"{sec}.{sub}")
        elif m.group(1) == "paragraph":
            if not in_subsec_paragraph_section.get(sec):
                sub += 1
                last = ("sub", f"{sec}.{sub}")
        elif m.group(2):
            env = m.group(2).rstrip("*")
            if env == "figure":
                fig += 1
                last = ("fig", str(fig))
            elif env == "table":
                tab += 1
                last = ("tab", str(tab))
            elif env == "algorithm":
                alg += 1
                last = ("alg", str(alg))
            else:
                eq += 1
                last = ("eq", str(eq))
        elif m.group(3):
            labels[m.group(3)] = last[1]
    return labels


def prepare_source():
    tex = (PAPER / "main.tex").read_text(encoding="utf-8")
    body = tex[tex.index("\\section{Introduction}"):tex.index("\\bibliographystyle")]
    body = strip_comments(body)
    # objectives list -> numbered equations
    s = body.index("\\begin{itemize}", body.index("\\label{sec:objectives}"))
    e = body.index("\\end{itemize}", s) + len("\\end{itemize}")
    lead = body.index("For a student with logits", body.index("\\label{sec:objectives}"))
    body = body[:s] + OBJECTIVES + body[e:]
    body = body.replace("logits $t$, class prior $\\pi$ and temperature $T$:",
                        "logits $t$, class prior $\\pi$, temperature $T$ and mini-batch $B$, the objectives are as follows.")
    assert lead < s
    return body


def build():
    b = Builder(TEMPLATE)
    src = prepare_source()
    b.labels = numbering_pass(src)

    # ---------------- title block (template paragraphs P1..P11, edited in place)
    P = b.title_block

    def set_text(p, segments, keep_runs=1):
        runs = p._p.findall(q("w:r"))
        tmpl = copy.deepcopy(runs[0]) if runs else None
        for r in runs:
            p._p.remove(r)
        for text, fmt in segments:
            r = copy.deepcopy(tmpl)
            rpr = r.find(q("w:rPr"))
            for e in rpr.findall(q("w:color")) + rpr.findall(q("w:vertAlign")):
                rpr.remove(e)
            if fmt.get("sup"):
                v = etree.SubElement(rpr, q("w:vertAlign"))
                v.set(q("w:val"), "superscript")
            if fmt.get("red"):
                c = etree.SubElement(rpr, q("w:color"))
                c.set(q("w:val"), "FF0000")
            for t in r.findall(q("w:t")):
                r.remove(t)
            for el in list(r):
                if el.tag != q("w:rPr"):
                    r.remove(el)
            t = etree.SubElement(r, q("w:t"))
            t.text = text
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            p._p.append(r)

    title = ("When Compression Forgets Rare Attacks: Knowledge Distillation and Quantization for "
             "Resource-Constrained IoT Intrusion Detection")
    set_text(P[1], [(title, {})])
    set_text(P[3], [("[First Name Family Name]", {"red": True}), ("1", {"sup": True, "red": True}),
                    ("*          ", {"red": True}), ("[First Name Family Name]", {"red": True}),
                    ("2", {"sup": True, "red": True})])
    set_text(P[5], [("1", {"sup": True}), ("[Author\u2019s Affiliation 1, Full Postal Address, Country]", {"red": True})])
    set_text(P[6], [("2", {"sup": True}), ("[Author\u2019s Affiliation 2, Full Postal Address, Country]", {"red": True})])
    set_text(P[7], [("* Corresponding author\u2019s Email: ", {}), ("[email]", {"red": True})])
    words = len(ABSTRACT.split())
    assert words <= 150, f"abstract has {words} words (IJIES limit 150)"
    abs_runs = P[10]._p.findall(q("w:r"))
    for r in abs_runs[1:]:
        P[10]._p.remove(r)
    r = copy.deepcopy(abs_runs[0])
    r.find(q("w:rPr")).remove(r.find(q("w:rPr")).find(q("w:b")))
    r.find(q("w:t")).text = ABSTRACT
    P[10]._p.append(r)
    kw_runs = P[11]._p.findall(q("w:r"))
    kw_runs[-1].find(q("w:t")).text = KEYWORDS
    for r in kw_runs[2:-1]:
        P[11]._p.remove(r)

    # ---------------- remove the template's instructional body
    body = b.body
    started = False
    for el in list(body):
        if el is P[14]._p:
            started = True
            continue
        if started and el.tag != q("w:sectPr"):
            body.remove(el)
    final = body.find(q("w:sectPr"))
    body.replace(final, copy.deepcopy(b.sect_two))

    # ---------------- main text
    pos, n = 0, len(src)
    section_has_sub = False
    in_section_star = False

    def text_blocks(chunk):
        for para in re.split(r"\n\s*\n", chunk):
            para = para.strip()
            if para:
                yield para

    env_re = re.compile(r"\\begin\{(figure\*?|table\*?|algorithm|enumerate|itemize|equation)\}(?:\[[^\]]*\])?")
    head_re = re.compile(r"\\(section|subsection|paragraph)(\*?)\{")

    def emit_text(chunk):
        for para in text_blocks(chunk):
            para = re.sub(r"\\label\{[^}]*\}", "", para).strip()
            if not para:
                continue
            key = "textbody" if in_section_star else "body"
            b.para(key, para)

    while pos < n:
        mh = head_re.search(src, pos)
        me = env_re.search(src, pos)
        nxt = min([m for m in (mh, me) if m], key=lambda m: m.start(), default=None)
        if nxt is None:
            emit_text(src[pos:])
            break
        emit_text(src[pos:nxt.start()])
        if nxt is mh:
            title_txt, after = find_group(src, mh.end() - 1)
            kind, star = mh.group(1), mh.group(2)
            if kind == "section" and star:
                in_section_star = True
                b.para("chapter", segments=[("t", "Data and Code Availability", {})])
            elif kind == "section":
                in_section_star = False
                section_has_sub = False
                b.heading(1, title_txt)
            elif kind == "subsection":
                section_has_sub = True
                b.heading(2, title_txt)
            else:   # \paragraph: second-order heading if the section has no subsections
                if not section_has_sub:
                    b.heading(2, title_txt)
                else:
                    # run-in bold lead of the following paragraph
                    rest = src[after:]
                    para_end = re.search(r"\n\s*\n", rest)
                    para_txt = rest[:para_end.start()] if para_end else rest
                    b.para("body", f"\\textbf{{{title_txt}.}} " + para_txt.strip())
                    pos = after + (para_end.end() if para_end else len(rest))
                    continue
            pos = after
            continue
        env = nxt.group(1)
        end_tag = f"\\end{{{env}}}"
        e = src.index(end_tag, nxt.end())
        inner = src[nxt.end():e]
        pos = e + len(end_tag)
        if env == "equation":
            lab = re.search(r"\\label\{([^}]+)\}", inner).group(1)
            math = re.sub(r"\\label\{[^}]*\}", "", inner).strip()
            b.equation(math, b.labels[lab])
        elif env in ("enumerate", "itemize"):
            items = [x.strip() for x in re.split(r"\\item\b", inner) if x.strip()]
            for k, it in enumerate(items, 1):
                marker = f"({k}) " if env == "enumerate" else "\u2013 "
                b.para("body", marker + it)
        elif env.startswith("figure"):
            wide = env.endswith("*")
            cap = re.search(r"\\caption\{", inner)
            caption, _ = find_group(inner, cap.end() - 1)
            lab = re.search(r"\\label\{([^}]+)\}", inner).group(1)
            g = re.search(r"\\(?:includegraphics(?:\[[^\]]*\])?\{([^}]+)\}|input\{figures/([^}]+)\})", inner)
            name = Path(g.group(1) or g.group(2)).stem
            img = PAPER / "figures" / f"{name}.png"
            if wide:
                b.section_break(two_columns_before=True)
            b.picture(img, FULL_CM if wide else COL_CM)
            b.para("caption", segments=[("t", f"Fig. {b.labels[lab]}. ", {})] + b.inline.parse(caption), size=10)
            b.append(b._clone("blank"))
            if wide:
                b.section_break(two_columns_before=False)
        elif env.startswith("table"):
            wide = env.endswith("*") or "tab:teachers" in inner
            cap = re.search(r"\\caption\{", inner)
            caption, _ = find_group(inner, cap.end() - 1)
            lab = re.search(r"\\label\{([^}]+)\}", inner).group(1)
            tpath = PAPER / (re.search(r"\\(?:fittable|input)\{([^}]+)\}", inner).group(1) + ".tex")
            if wide:
                b.section_break(two_columns_before=True)
            b.para("tabtitle", segments=[("t", f"Table {b.labels[lab]}. ", {})] + b.inline.parse(caption), size=10)
            b.table(tpath, FULL_W if wide else COL_W)
            b.append(b._clone("blank"))
            if wide:
                b.section_break(two_columns_before=False)
        elif env == "algorithm":
            cap = re.search(r"\\caption\{", inner)
            caption, _ = find_group(inner, cap.end() - 1)
            lab = re.search(r"\\label\{([^}]+)\}", inner).group(1)
            alg = inner[inner.index("\\begin{algorithmic}"):inner.index("\\end{algorithmic}")]
            alg = re.sub(r"\\begin\{algorithmic\}(\[\d+\])?", "", alg)
            b.algorithm(caption, b.labels[lab], alg)

    # ---------------- closing sections (template order)
    for name, text in CLOSING.items():
        if text is None:
            continue
        b.para("chapter", segments=[("t", name, {})])
        b.para("textbody", text)

    # ---------------- references (order of first citation)
    b.para("refhead", segments=[("t", "References", {})])
    unused = set(b.bib) - set(b.cite_order)
    if unused:
        raise SystemExit(f"references never cited: {sorted(unused)}")
    for key in sorted(b.cite_order, key=b.cite_order.get):
        kind, f = b.bib[key]
        p = b._clone("refitem")
        # the template's 346-twip hanging indent fits "[9]" but not "[10]"-"[43]"
        ind = p.find(q("w:pPr") + "/" + q("w:ind"))
        ind.set(q("w:left"), "454")
        ind.set(q("w:hanging"), "454")
        b.add_runs(p, [("t", t, fmt) for t, fmt in reference_runs(kind, f)], size=11)
        b.append(p)

    b.resolve_math()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    b.doc.save(OUT)
    print(f"wrote {OUT}  ({len(b.cite_order)} references, abstract {len(ABSTRACT.split())} words)")


if __name__ == "__main__":
    build()
